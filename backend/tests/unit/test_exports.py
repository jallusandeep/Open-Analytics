import csv
from datetime import date
from io import BytesIO, StringIO

import duckdb
from fastapi import HTTPException
from openpyxl import load_workbook
import pytest

from app.api.v1 import data_export_routes as exports

pytestmark = pytest.mark.unit


def test_filtered_export_only_contains_supplied_matching_rows():
    response = exports.download_filtered_dataset("equity_news", exports.FilteredExport(
        headers=["Title", "Score"], rows=[["Matching article", 12.34]]
    ))
    assert list(csv.reader(StringIO(response.body.decode("utf-8-sig")))) == [
        ["Title", "Score"], ["Matching article", "12.34"]
    ]


def test_filtered_excel_preserves_numeric_values_and_formula_safety():
    response = exports.download_filtered_dataset("equity_news", exports.FilteredExport(
        headers=["Title", "Score"], rows=[["=unsafe", 12.34]]
    ), format="xlsx")
    sheet = load_workbook(BytesIO(response.body)).active
    assert sheet["A2"].data_type == "s"
    assert sheet["B2"].value == 12.34


def test_filtered_export_rejects_mismatched_columns():
    with pytest.raises(HTTPException) as error:
        exports.download_filtered_dataset("equity_news", exports.FilteredExport(
            headers=["Title"], rows=[["Article", 1]]
        ))
    assert error.value.status_code == 400


@pytest.fixture
def export_db(tmp_path, monkeypatch):
    path = str(tmp_path / "export.duckdb")
    conn = duckdb.connect(path)
    conn.execute("CREATE TABLE equity_news (title VARCHAR, published_at TIMESTAMP, score DOUBLE)")
    conn.execute("INSERT INTO equity_news VALUES ('=formula', '2026-09-01 12:00:00', 1.2345), ('last', '2026-09-02 23:59:59', 2.0), ('later', '2026-09-03 00:00:00', 3.0)")
    conn.close()
    monkeypatch.setattr(exports, "get_connection", lambda: duckdb.connect(path))


def test_csv_date_range_is_inclusive_and_escapes_formulas(export_db):
    response = exports.download_dataset("equity_news", start_date=date(2026, 9, 1), end_date=date(2026, 9, 2))
    rows = list(csv.reader(StringIO(response.body.decode("utf-8-sig"))))
    assert len(rows) == 3
    assert rows[1][0] == "'=formula"
    assert rows[2][0] == "last"


def test_excel_preserves_numbers_and_formula_text(export_db):
    response = exports.download_dataset("equity_news", format="xlsx")
    workbook = load_workbook(BytesIO(response.body))
    sheet = workbook.active
    assert sheet["A2"].value == "=formula"
    assert sheet["A2"].data_type == "s"
    assert sheet["C2"].value == 1.2345


def test_date_metadata(export_db):
    assert exports.export_options("equity_news")["default_date_column"] == "published_at"


@pytest.mark.parametrize("kwargs", [{"date_column": "invalid"}, {"start_date": date(2026, 9, 3), "end_date": date(2026, 9, 1)}])
def test_invalid_range_and_column_rejected(export_db, kwargs):
    with pytest.raises(HTTPException) as error:
        exports.download_dataset("equity_news", **kwargs)
    assert error.value.status_code == 400


def test_export_limit_never_silently_truncates(export_db, monkeypatch):
    monkeypatch.setattr(exports, "MAX_ROWS", 1)
    with pytest.raises(HTTPException) as error:
        exports.download_dataset("equity_news")
    assert error.value.status_code == 413
